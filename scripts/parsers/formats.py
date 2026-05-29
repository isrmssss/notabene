"""
Парсеры для различных форматов файлов

Поддерживает: PDF, DOCX, PPTX, TXT, MD, CSV, EPUB
"""

import csv
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from ebooklib import epub
import ebooklib
from bs4 import BeautifulSoup

from parsers.base import BaseParser
from parsers.utils import clean_text


class PlainTextParser(BaseParser):
    """Парсер для текстовых файлов"""
    
    def parse(self, file_path: Path) -> str:
        encodings = ["utf-8", "windows-1251", "cp1252"]
        
        for encoding in encodings:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    return clean_text(f.read())
            except UnicodeDecodeError:
                continue
        
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return clean_text(f.read())


class PdfParser(BaseParser):
    """Парсер для PDF файлов с поддержкой OCR для изображений"""
    
    def __init__(self, ocr_parser):
        self.ocr_parser = ocr_parser
    
    def parse(self, file_path: Path) -> str:
        text = []
        reader = PdfReader(file_path)
        
        for page_idx, page in enumerate(reader.pages):
            try:
                extracted = page.extract_text()
                if extracted and extracted.strip():
                    text.append(extracted)
            except Exception:
                pass
            
            try:
                if hasattr(page, 'images'):
                    for image_file_object in page.images:
                        ocr_text = self.ocr_parser.parse_bytes(image_file_object.data)
                        if ocr_text:
                            text.append(ocr_text)
            except Exception:
                pass
        
        return clean_text("\n".join(text))


class DocxParser(BaseParser):
    """Парсер для DOCX файлов с поддержкой встроенных изображений"""
    
    def __init__(self, ocr_parser):
        self.ocr_parser = ocr_parser
    
    def parse(self, file_path: Path) -> str:
        text = []
        doc = Document(file_path)
        
        # Параграфы с встроенными изображениями
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)
            
            for run in paragraph.runs:
                inline_shapes = run._element.xpath('.//wp:inline | .//wp:anchor')
                for shape_elem in inline_shapes:
                    try:
                        blips = shape_elem.xpath('.//a:blip')
                        for blip in blips:
                            embed_id = blip.get(
                                '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
                            )
                            if embed_id and embed_id in doc.part.related_parts:
                                part = doc.part.related_parts[embed_id]
                                if "image" in part.content_type:
                                    ocr_text = self.ocr_parser.parse_bytes(part.blob)
                                    if ocr_text:
                                        text.append(ocr_text)
                    except Exception:
                        pass
        
        # Таблицы
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if paragraph.text.strip():
                            text.append(paragraph.text)
        
        return clean_text("\n".join(text))


class PptxParser(BaseParser):
    """Парсер для PPTX файлов с поддержкой встроенных изображений"""
    
    def __init__(self, ocr_parser):
        self.ocr_parser = ocr_parser
    
    def parse(self, file_path: Path) -> str:
        text = []
        prs = Presentation(file_path)
        
        for slide_idx, slide in enumerate(prs.slides):
            # Текст в shapes
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text.append(shape.text)
            
            # Встроенные изображения
            for shape in slide.shapes:
                try:
                    if hasattr(shape, "image"):
                        image = shape.image
                        ocr_text = self.ocr_parser.parse_bytes(image.blob)
                        if ocr_text:
                            text.append(ocr_text)
                except Exception:
                    pass
            
            # Изображения в связанных ресурсах
            try:
                for rel_id, part in slide.part.related_parts.items():
                    if "image" in part.content_type:
                        ocr_text = self.ocr_parser.parse_bytes(part.blob)
                        if ocr_text:
                            text.append(ocr_text)
            except Exception:
                pass
        
        return clean_text("\n".join(text))


class CsvParser(BaseParser):
    """Парсер для CSV файлов"""
    
    def parse(self, file_path: Path) -> str:
        text = []
        
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            for row in reader:
                text.append(" ".join(row))
        
        return clean_text("\n".join(text))


class EpubParser(BaseParser):
    """Парсер для EPUB файлов"""
    
    def parse(self, file_path: Path) -> str:
        text = []
        
        book = epub.read_epub(file_path)
        for item in book.get_items():
            if item.get_type() == ebooklib.ITEM_DOCUMENT:
                soup = BeautifulSoup(item.get_content(), "html.parser")
                text.append(soup.get_text())
        
        return clean_text("\n".join(text))
