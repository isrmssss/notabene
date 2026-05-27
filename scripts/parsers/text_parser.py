import csv
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from ebooklib import epub
import ebooklib
from bs4 import BeautifulSoup
from parsers.base import BaseParser
from parsers.ocr_parser import OcrParser
from parsers.utils import clean_text

class PlainTextParser(BaseParser):
    def parse(self, file_path: Path) -> str:
        encodings = ["utf-8", "windows-1251", "cp1252"]
        for encoding in encodings:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    return clean_text(f.read())
            except UnicodeDecodeError:
                continue
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return clean_text(f.read())
        except Exception:
            return ""

class PdfParser(BaseParser):
    def __init__(self, ocr_parser: OcrParser):
        self.ocr_parser = ocr_parser

    def parse(self, file_path: Path) -> str:
        text = []
        try:
            reader = PdfReader(file_path)
            for page_idx, page in enumerate(reader.pages):
                try:
                    # Пытаемся извлечь обычный текст
                    extracted = page.extract_text()
                    if extracted and extracted.strip():
                        text.append(extracted)
                except Exception as e:
                    print(f"PDF text extraction error on page {page_idx}: {e}")
                
                # Обрабатываем изображения на странице
                try:
                    # Метод 1: Через page.images
                    if hasattr(page, 'images'):
                        for image_file_object in page.images:
                            try:
                                ocr_text = self.ocr_parser.parse_bytes(image_file_object.data)
                                if ocr_text:
                                    text.append(ocr_text)
                            except Exception:
                                pass
                except Exception as e:
                    print(f"PDF image extraction error on page {page_idx}: {e}")
        except Exception as e:
            print(f"PdfParser Error: {e}")
            return ""
        return clean_text("\n".join(text))

class DocxParser(BaseParser):
    def __init__(self, ocr_parser: OcrParser):
        self.ocr_parser = ocr_parser

    def parse(self, file_path: Path) -> str:
        text = []
        try:
            doc = Document(file_path)
            
            # Обрабатываем параграфы в порядке с встроенными изображениями
            for paragraph in doc.paragraphs:
                # Сначала добавляем текст параграфа
                if paragraph.text.strip():
                    text.append(paragraph.text)
                
                # Затем ищем встроенные изображения в параграфе
                for run in paragraph.runs:
                    # Ищем inline и anchor элементы (встроенные рисунки)
                    inline_shapes = run._element.xpath('.//wp:inline | .//wp:anchor')
                    for shape_elem in inline_shapes:
                        try:
                            # Ищем blip элемент (ссылка на изображение)
                            blips = shape_elem.xpath('.//a:blip')
                            for blip in blips:
                                # Получаем ID встроенного ресурса
                                embed_id = blip.get(
                                    '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
                                )
                                if embed_id and embed_id in doc.part.related_parts:
                                    part = doc.part.related_parts[embed_id]
                                    if "image" in part.content_type:
                                        ocr_text = self.ocr_parser.parse_bytes(part.blob)
                                        if ocr_text:
                                            text.append(ocr_text)
                        except Exception:
                            pass
            
            # Обрабатываем таблицы (если есть)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            if paragraph.text.strip():
                                text.append(paragraph.text)
                            
                            # Ищем изображения в таблице
                            for run in paragraph.runs:
                                inline_shapes = run._element.xpath('.//wp:inline | .//wp:anchor')
                                for shape_elem in inline_shapes:
                                    try:
                                        blips = shape_elem.xpath('.//a:blip')
                                        for blip in blips:
                                            embed_id = blip.get(
                                                '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
                                            )
                                            if embed_id and embed_id in doc.part.related_parts:
                                                part = doc.part.related_parts[embed_id]
                                                if "image" in part.content_type:
                                                    ocr_text = self.ocr_parser.parse_bytes(part.blob)
                                                    if ocr_text:
                                                        text.append(ocr_text)
                                    except Exception:
                                        pass
        except Exception as e:
            print(f"DocxParser Error: {e}")
            return ""
        return clean_text("\n".join(text))

class PptxParser(BaseParser):
    def __init__(self, ocr_parser: OcrParser):
        self.ocr_parser = ocr_parser

    def parse(self, file_path: Path) -> str:
        text = []
        try:
            prs = Presentation(file_path)
            for slide_idx, slide in enumerate(prs.slides):
                # Обрабатываем текст в shapes
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text.append(shape.text)
                
                # Обрабатываем встроенные изображения в shapes
                for shape in slide.shapes:
                    try:
                        if hasattr(shape, "image"):
                            # Shape содержит изображение
                            image = shape.image
                            ocr_text = self.ocr_parser.parse_bytes(image.blob)
                            if ocr_text:
                                text.append(ocr_text)
                    except Exception:
                        pass
                
                # Обрабатываем изображения в связанных ресурсах слайда
                try:
                    for rel_id, part in slide.part.related_parts.items():
                        try:
                            if "image" in part.content_type:
                                ocr_text = self.ocr_parser.parse_bytes(part.blob)
                                if ocr_text:
                                    text.append(ocr_text)
                        except Exception:
                            pass
                except Exception:
                    pass
        except Exception as e:
            print(f"PptxParser Error: {e}")
            return ""
        return clean_text("\n".join(text))

class CsvParser(BaseParser):
    def parse(self, file_path: Path) -> str:
        text = []
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                for row in reader:
                    text.append(" ".join(row))
        except Exception:
            return ""
        return clean_text("\n".join(text))

class EpubParser(BaseParser):
    def parse(self, file_path: Path) -> str:
        text = []
        try:
            book = epub.read_epub(file_path)
            for item in book.get_items():
                if item.get_type() == ebooklib.ITEM_DOCUMENT:
                    soup = BeautifulSoup(item.get_content(), "html.parser")
                    text.append(soup.get_text())
        except Exception:
            return ""
        return clean_text("\n".join(text))